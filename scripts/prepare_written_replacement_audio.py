from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from pypinyin import Style, lazy_pinyin
import pycantonese


WORKBOOK_PATH = Path("/home/ubuntu/upload/繁體認字樂_1260詞_建議改成_書面語版.xlsx")
TERMS_PATH = Path("/home/ubuntu/hong-kong-chinese-word-library/client/src/data/termsV2.json")
AUDIO_PATH = Path("/home/ubuntu/hong-kong-chinese-word-library/client/src/data/topicAudioV2.ts")
OUTPUT_PATH = Path("/home/ubuntu/written_replacement_audio_plan.json")
BATCH_SIZE = 18


def normalized(value: object) -> str:
    return str(value or "").strip()


def jyutping_for(term: str) -> str:
    segments = pycantonese.characters_to_jyutping(term)
    readings = [reading for _, reading in segments if reading]
    return " ".join(readings)


def pinyin_for(term: str) -> str:
    return " ".join(lazy_pinyin(term, style=Style.TONE))


def split_batches(items: list[str]) -> list[list[str]]:
    return [items[index : index + BATCH_SIZE] for index in range(0, len(items), BATCH_SIZE)]


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    worksheet = workbook["完整詞庫1260詞"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [normalized(value) for value in next(rows)]
    indexes = {header: index for index, header in enumerate(headers)}
    terms = json.loads(TERMS_PATH.read_text(encoding="utf-8"))
    current_by_term = {item["term"]: item for item in terms}
    current_by_position = {
        (item["level"], item["mapId"], item["stage"], item["stageOrder"]): item
        for item in terms
    }
    audio_keys = set(
        re.findall(r'^\s*"(.+)": \{ cantonese:', AUDIO_PATH.read_text(encoding="utf-8"), re.MULTILINE)
    )

    replacements = []
    for row_number, row in enumerate(rows, start=2):
        old_term = normalized(row[indexes["詞語"]])
        suggested_term = normalized(row[indexes["建議改成"]])
        if not suggested_term or suggested_term == old_term:
            continue
        new_term = "團拜" if old_term == "互相" and suggested_term == "團聚" else suggested_term
        location_key = (
            int(row[indexes["第幾級"]]),
            normalized(row[indexes["地圖編號"]]),
            int(row[indexes["第幾關"]]),
            int(row[indexes["關內次序"]]),
        )
        existing = current_by_position.get(location_key) or current_by_term[old_term]
        replacements.append(
            {
                "row": row_number,
                "old": old_term,
                "new": new_term,
                "mapId": existing["mapId"],
                "stage": existing["stage"],
                "stageOrder": existing["stageOrder"],
                "reuseExistingAudio": new_term in audio_keys,
                "jyutping": jyutping_for(new_term),
                "pinyin": pinyin_for(new_term),
            }
        )

    new_audio_terms = sorted({item["new"] for item in replacements if not item["reuseExistingAudio"]})
    reusable_terms = sorted({item["new"] for item in replacements if item["reuseExistingAudio"]})
    output = {
        "replacementCount": len(replacements),
        "replacements": replacements,
        "newAudioTerms": new_audio_terms,
        "reusableExistingAudioTerms": reusable_terms,
        "cantoneseBatches": split_batches(new_audio_terms),
        "mandarinBatches": split_batches(new_audio_terms),
        "emptyJyutpingTerms": [item["new"] for item in replacements if not item["jyutping"]],
    }
    OUTPUT_PATH.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "replacementCount": output["replacementCount"],
        "newAudioTermCount": len(new_audio_terms),
        "reusableExistingAudioTermCount": len(reusable_terms),
        "emptyJyutpingTerms": output["emptyJyutpingTerms"],
        "cantoneseBatchCount": len(output["cantoneseBatches"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
