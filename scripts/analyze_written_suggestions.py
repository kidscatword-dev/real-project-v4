from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_PATH = Path("/home/ubuntu/upload/繁體認字樂_1260詞_建議改成_書面語版.xlsx")
TERMS_PATH = Path("/home/ubuntu/hong-kong-chinese-word-library/client/src/data/termsV2.json")
AUDIO_PATH = Path("/home/ubuntu/hong-kong-chinese-word-library/client/src/data/topicAudioV2.ts")
OUTPUT_PATH = Path("/home/ubuntu/written_suggestion_analysis.json")


def normalized(value: object) -> str:
    return str(value or "").strip()


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH, data_only=True, read_only=True)
    worksheet = workbook["完整詞庫1260詞"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [normalized(value) for value in next(rows)]
    indexes = {header: index for index, header in enumerate(headers)}

    suggestions = []
    for row_number, row in enumerate(rows, start=2):
        old_term = normalized(row[indexes["詞語"]])
        new_term = normalized(row[indexes["建議改成"]])
        if not new_term or new_term == old_term:
            continue
        suggestions.append(
            {
                "row": row_number,
                "old": old_term,
                "new": new_term,
                "level": int(row[indexes["第幾級"]]),
                "mapId": normalized(row[indexes["地圖編號"]]),
                "topic": normalized(row[indexes["地圖主題"]]),
                "stage": int(row[indexes["第幾關"]]),
                "position": int(row[indexes["關內次序"]]),
                "currentJyutping": normalized(row[indexes["粵拼"]]),
                "currentPinyin": normalized(row[indexes["普通話拼音"]]),
                "audioStatus": normalized(row[indexes["固定音檔狀態"]]),
            }
        )

    with TERMS_PATH.open(encoding="utf-8") as terms_file:
        terms = json.load(terms_file)
    existing_terms = {normalized(item.get("term")) for item in terms}
    audio_keys = set(re.findall(r'^\s*"(.+)": \{ cantonese:', AUDIO_PATH.read_text(encoding="utf-8"), re.MULTILINE))
    replacement_old_terms = {item["old"] for item in suggestions}
    unchanged_terms = existing_terms - replacement_old_terms
    suggested_counts = Counter(item["new"] for item in suggestions)
    suggestions_by_map = defaultdict(list)
    for item in suggestions:
        item["conflictsWithUnchangedTerm"] = item["new"] in unchanged_terms
        item["duplicateSuggestedDestination"] = suggested_counts[item["new"]] > 1
        item["hasExistingAudioForNewTerm"] = item["new"] in audio_keys
        suggestions_by_map[item["mapId"]].append(item)

    output = {
        "suggestionCount": len(suggestions),
        "mapsAffected": len(suggestions_by_map),
        "suggestions": suggestions,
        "duplicateSuggestedDestinations": sorted(
            term for term, count in suggested_counts.items() if count > 1
        ),
        "conflictsWithUnchangedTerms": sorted(
            item["new"] for item in suggestions if item["conflictsWithUnchangedTerm"]
        ),
        "reusableExistingAudio": sorted(
            item["new"] for item in suggestions if item["hasExistingAudioForNewTerm"]
        ),
        "newAudioRequired": sorted(
            item["new"] for item in suggestions if not item["hasExistingAudioForNewTerm"]
        ),
        "byMap": {
            map_id: {
                "topic": entries[0]["topic"],
                "count": len(entries),
                "replacements": [{"old": entry["old"], "new": entry["new"]} for entry in entries],
            }
            for map_id, entries in sorted(suggestions_by_map.items())
        },
    }
    OUTPUT_PATH.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({
        "suggestionCount": output["suggestionCount"],
        "mapsAffected": output["mapsAffected"],
        "duplicateSuggestedDestinations": output["duplicateSuggestedDestinations"],
        "conflictsWithUnchangedTerms": output["conflictsWithUnchangedTerms"],
        "reusableExistingAudioCount": len(output["reusableExistingAudio"]),
        "newAudioRequiredCount": len(output["newAudioRequired"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
