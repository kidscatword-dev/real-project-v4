from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parent.parent
TERMS_PATH = PROJECT_ROOT / "client/src/data/termsV2.json"
AUDIO_PATH = PROJECT_ROOT / "client/src/data/topicAudioV2.ts"
OUTPUT_PATH = PROJECT_ROOT / "docs/繁體認字樂_粵普讀音核對總表.xlsx"
SITE_ORIGIN = "https://hkchineselib-mcq79f2x.manus.space"
REVIEW_SOURCES = {"漏字.xlsx 推薦加入", "本輪補充詞（家長確認採用）"}


def load_audio_mapping() -> dict[str, dict[str, str]]:
    source = AUDIO_PATH.read_text(encoding="utf-8")
    pattern = re.compile(
        r'^\s*"([^"]+)": \{ cantonese: "([^"]+)", mandarin: "([^"]+)" \},$',
        re.MULTILINE,
    )
    return {
        term: {"cantonese": cantonese, "mandarin": mandarin}
        for term, cantonese, mandarin in pattern.findall(source)
    }


def absolute_audio_url(path: str) -> str:
    return f"{SITE_ORIGIN}{path}"


def build_workbook(terms: list[dict], audio_map: dict[str, dict[str, str]]) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "使用說明"
    review = workbook.create_sheet("讀音核對總表")

    summary.sheet_view.showGridLines = False
    summary["A1"] = "繁體認字樂｜粵語及普通話讀音核對總表"
    summary["A1"].font = Font(name="Microsoft JhengHei", size=18, bold=True, color="1E3A5F")
    summary.merge_cells("A1:F1")
    summary["A3"] = "如何使用"
    summary["A3"].font = Font(name="Microsoft JhengHei", size=13, bold=True, color="1E3A5F")
    instructions = [
        "在「讀音核對總表」按「▶ 粵語」或「▶ 普通話」，便可直接播放該詞的固定讀音。",
        "聽完後，在兩個「核對結果」欄選擇「正常」、「無聲」、「讀錯」、「音量太小」或「其他」。",
        "如有問題，請在「問題／建議正確讀法」寫下你聽到的情況或正確讀法；例如：粵語「啤」字讀錯。",
        "完成後把這個 Excel 檔交回即可；我會只修改你標記的詞語，不會猜測性更換其他讀音。",
    ]
    for row, text in enumerate(instructions, start=4):
        summary.cell(row, 1, f"{row - 3}. {text}")
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        summary.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

    summary["A10"] = "總覽"
    summary["A10"].font = Font(name="Microsoft JhengHei", size=13, bold=True, color="1E3A5F")
    overview = [
        ("正式詞語", len(terms)),
        ("固定讀音段數", len(terms) * 2),
        ("後加詞語", sum(item["source"] in REVIEW_SOURCES for item in terms)),
        ("原有詞語", sum(item["source"] not in REVIEW_SOURCES for item in terms)),
        ("粵語已標記正常", "=COUNTIF('讀音核對總表'!K:K,\"正常\")"),
        ("普通話已標記正常", "=COUNTIF('讀音核對總表'!L:L,\"正常\")"),
        ("待處理讀音", "=COUNTIF('讀音核對總表'!K:K,\"無聲\")+COUNTIF('讀音核對總表'!K:K,\"讀錯\")+COUNTIF('讀音核對總表'!K:K,\"音量太小\")+COUNTIF('讀音核對總表'!L:L,\"無聲\")+COUNTIF('讀音核對總表'!L:L,\"讀錯\")+COUNTIF('讀音核對總表'!L:L,\"音量太小\")"),
    ]
    for row, (label, value) in enumerate(overview, start=11):
        summary.cell(row, 1, label)
        summary.cell(row, 2, value)
        summary.cell(row, 1).font = Font(name="Microsoft JhengHei", bold=True)

    source_counts = Counter(item["source"] for item in terms)
    summary["D10"] = "詞條來源"
    summary["D10"].font = Font(name="Microsoft JhengHei", size=13, bold=True, color="1E3A5F")
    for row, (label, count) in enumerate(sorted(source_counts.items()), start=11):
        summary.cell(row, 4, label)
        summary.cell(row, 5, count)

    for column, width in {"A": 74, "B": 18, "C": 4, "D": 30, "E": 12, "F": 4}.items():
        summary.column_dimensions[column].width = width
    for row in range(4, 8):
        summary.row_dimensions[row].height = 34

    headers = [
        "序號", "級別", "地圖", "主題", "關卡", "關內次序", "字詞",
        "粵語試聽", "普通話試聽", "詞條來源", "粵語核對結果", "普通話核對結果",
        "問題／建議正確讀法", "已回報給我",
    ]
    header_fill = PatternFill("solid", fgColor="1E6091")
    header_font = Font(name="Microsoft JhengHei", bold=True, color="FFFFFF")
    thin_blue = Side(style="thin", color="C8D9E8")
    for column, title in enumerate(headers, start=1):
        cell = review.cell(1, column, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_blue)
    review.row_dimensions[1].height = 32

    for index, item in enumerate(terms, start=1):
        row = index + 1
        audio = audio_map[item["term"]]
        values = [
            index,
            f"第 {item['level']} 級",
            item["mapId"],
            item["topic"],
            f"第 {item['stage']} 關",
            item["stageOrder"],
            item["term"],
            "▶ 粵語",
            "▶ 普通話",
            item["source"],
            "未檢查",
            "未檢查",
            "",
            "",
        ]
        for column, value in enumerate(values, start=1):
            cell = review.cell(row, column, value)
            cell.font = Font(name="Microsoft JhengHei", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=column in {4, 10, 13})
            cell.border = Border(bottom=Side(style="hair", color="E8EEF4"))
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FBFE")
        for column, language in [(8, "cantonese"), (9, "mandarin")]:
            cell = review.cell(row, column)
            cell.hyperlink = absolute_audio_url(audio[language])
            cell.style = "Hyperlink"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    last_row = len(terms) + 1
    review.freeze_panes = "G2"
    review.auto_filter.ref = f"A1:N{last_row}"
    review.sheet_view.zoomScale = 85
    widths = {
        "A": 8, "B": 10, "C": 9, "D": 25, "E": 10, "F": 11, "G": 16,
        "H": 12, "I": 13, "J": 29, "K": 15, "L": 16, "M": 38, "N": 14,
    }
    for column, width in widths.items():
        review.column_dimensions[column].width = width
    for row in range(2, last_row + 1):
        review.row_dimensions[row].height = 25

    choices = '"未檢查,正常,無聲,讀錯,音量太小,其他"'
    result_validation = DataValidation(type="list", formula1=choices, allow_blank=False)
    result_validation.error = "請從清單選擇核對結果。"
    result_validation.errorTitle = "核對結果"
    review.add_data_validation(result_validation)
    result_validation.add(f"K2:L{last_row}")
    reported_validation = DataValidation(type="list", formula1='"未回報,已回報"', allow_blank=True)
    review.add_data_validation(reported_validation)
    reported_validation.add(f"N2:N{last_row}")

    red_fill = PatternFill("solid", fgColor="FDE2E2")
    orange_fill = PatternFill("solid", fgColor="FFF0D6")
    green_fill = PatternFill("solid", fgColor="E5F3E7")
    review.conditional_formatting.add(f"K2:L{last_row}", FormulaRule(formula=['K2="讀錯"'], fill=red_fill))
    review.conditional_formatting.add(f"K2:L{last_row}", FormulaRule(formula=['K2="無聲"'], fill=red_fill))
    review.conditional_formatting.add(f"K2:L{last_row}", FormulaRule(formula=['K2="音量太小"'], fill=orange_fill))
    review.conditional_formatting.add(f"K2:L{last_row}", FormulaRule(formula=['K2="正常"'], fill=green_fill))

    return workbook


def main() -> None:
    terms = json.loads(TERMS_PATH.read_text(encoding="utf-8"))
    audio_map = load_audio_mapping()
    if len(terms) != 1050:
        raise ValueError(f"預期 1,050 個詞，實際為 {len(terms)}。")
    missing = [item["term"] for item in terms if item["term"] not in audio_map]
    if missing:
        raise ValueError(f"缺少音檔映射：{'、'.join(missing[:10])}")
    workbook = build_workbook(terms, audio_map)
    workbook.save(OUTPUT_PATH)
    print(f"已輸出 {len(terms)} 個詞、{len(terms) * 2} 段讀音：{OUTPUT_PATH}")


if __name__ == "__main__":
    main()
