from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_PATH = Path(__file__).resolve().parent.parent / "docs/繁體認字樂_粵普讀音核對總表.xlsx"
EXPECTED_HEADERS = [
    "序號", "級別", "地圖", "主題", "關卡", "關內次序", "字詞",
    "粵語試聽", "普通話試聽", "詞條來源", "粵語核對結果", "普通話核對結果",
    "問題／建議正確讀法", "已回報給我",
]


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH, data_only=False)
    assert workbook.sheetnames == ["使用說明", "讀音核對總表"]
    review = workbook["讀音核對總表"]
    assert review.max_row == 1051, f"詞語列數錯誤：{review.max_row - 1}"
    assert review.max_column == 14, f"欄位數錯誤：{review.max_column}"
    assert [cell.value for cell in review[1]] == EXPECTED_HEADERS
    assert review.freeze_panes == "G2"
    assert review.auto_filter.ref == "A1:N1051"

    terms = set()
    for row in range(2, 1052):
        terms.add(review.cell(row, 7).value)
        assert review.cell(row, 8).value == "▶ 粵語"
        assert review.cell(row, 9).value == "▶ 普通話"
        assert review.cell(row, 8).hyperlink and review.cell(row, 8).hyperlink.target.startswith("https://hkchineselib-mcq79f2x.manus.space/manus-storage/")
        assert review.cell(row, 9).hyperlink and review.cell(row, 9).hyperlink.target.startswith("https://hkchineselib-mcq79f2x.manus.space/manus-storage/")
        assert review.cell(row, 11).value == "未檢查"
        assert review.cell(row, 12).value == "未檢查"
    assert len(terms) == 1050, f"詞語重複或缺漏：{len(terms)}"
    assert len(review.data_validations.dataValidation) == 2
    print("核對總表驗證通過：1,050 個不重複詞、2,100 條雙語試聽連結及填寫欄位均完整。")


if __name__ == "__main__":
    main()
