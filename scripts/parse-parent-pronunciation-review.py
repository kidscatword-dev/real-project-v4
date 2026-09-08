from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
INPUT_PATH = Path("/home/ubuntu/upload/繁體認字樂_粵普讀音核對總表.xlsx")
OUTPUT_PATH = PROJECT_ROOT / "docs/家長回傳_讀音問題清單.json"
SUMMARY_PATH = PROJECT_ROOT / "docs/家長回傳_讀音問題摘要.md"
REVIEW_SHEET = "讀音核對總表"
ISSUE_STATUSES = {"無聲", "讀錯", "音量太小", "其他"}


def normalise(value: object) -> str:
    return str(value or "").strip()


def main() -> None:
    if not INPUT_PATH.exists():
        raise FileNotFoundError(f"找不到家長回傳的核對表：{INPUT_PATH}")

    workbook = load_workbook(INPUT_PATH, data_only=False)
    if REVIEW_SHEET not in workbook.sheetnames:
        raise ValueError(f"核對表缺少工作表：{REVIEW_SHEET}")
    sheet = workbook[REVIEW_SHEET]
    header = {normalise(cell.value): cell.column for cell in sheet[1]}
    required = {"字詞", "粵語核對結果", "普通話核對結果", "問題／建議正確讀法"}
    missing_headers = required - set(header)
    if missing_headers:
        raise ValueError(f"核對表缺少欄位：{'、'.join(sorted(missing_headers))}")

    issues: list[dict[str, object]] = []
    statuses: Counter[str] = Counter()
    checked_count = 0
    for row in range(2, sheet.max_row + 1):
        term = normalise(sheet.cell(row, header["字詞"]).value)
        if not term:
            continue
        note = normalise(sheet.cell(row, header["問題／建議正確讀法"]).value)
        base = {
            "row": row,
            "term": term,
            "level": normalise(sheet.cell(row, header.get("級別", 0)).value),
            "mapId": normalise(sheet.cell(row, header.get("地圖", 0)).value),
            "topic": normalise(sheet.cell(row, header.get("主題", 0)).value),
            "stage": normalise(sheet.cell(row, header.get("關卡", 0)).value),
            "source": normalise(sheet.cell(row, header.get("詞條來源", 0)).value),
            "note": note,
        }
        for language, status_column, listen_column in (
            ("cantonese", "粵語核對結果", "粵語試聽"),
            ("mandarin", "普通話核對結果", "普通話試聽"),
        ):
            status = normalise(sheet.cell(row, header[status_column]).value)
            if status != "未檢查":
                checked_count += 1
            if status not in ISSUE_STATUSES:
                continue
            listen_cell = sheet.cell(row, header[listen_column])
            issues.append({
                **base,
                "language": language,
                "status": status,
                "audioUrl": listen_cell.hyperlink.target if listen_cell.hyperlink else "",
            })
            statuses[status] += 1

    payload = {
        "input": str(INPUT_PATH),
        "totalTerms": sheet.max_row - 1,
        "checkedAudioCells": checked_count,
        "issues": issues,
        "issueCounts": dict(sorted(statuses.items())),
    }
    OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# 家長回傳讀音問題摘要",
        "",
        f"- 詞語列數：{payload['totalTerms']}",
        f"- 已填寫核對結果的讀音欄：{checked_count}",
        f"- 需處理讀音：{len(issues)}",
        "",
        "| 問題類型 | 數量 |",
        "| --- | ---: |",
    ]
    for status, count in sorted(statuses.items()):
        lines.append(f"| {status} | {count} |")
    lines.extend(["", "| 詞語 | 語言 | 問題 | 地圖 | 備註 |", "| --- | --- | --- | --- | --- |"])
    for item in issues:
        language_label = "粵語" if item["language"] == "cantonese" else "普通話"
        lines.append(f"| {item['term']} | {language_label} | {item['status']} | {item['mapId']} | {item['note']} |")
    SUMMARY_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"checkedAudioCells": checked_count, "issueCounts": dict(statuses), "issues": len(issues)}, ensure_ascii=False))
    print(f"已輸出：{OUTPUT_PATH}")
    print(f"已輸出：{SUMMARY_PATH}")


if __name__ == "__main__":
    main()
